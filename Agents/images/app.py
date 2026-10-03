import streamlit as st
from docx import Document
from graph import app as agent_system

st.set_page_config(page_title="Multi-Agent System", layout="wide")
st.title("🤖 Multi-Agent Document Generation Pipeline")

uploaded_file = st.file_uploader("Client PDF Upload Karein", type=["pdf"])

if uploaded_file and st.button("Start Agents Processing 🚀"):
    pdf_path = "temp_uploaded.pdf"
    with open(pdf_path, "wb") as f:
        f.write(uploaded_file.getbuffer())

    with st.spinner("Agents processing (Security -> Intake -> Requirements -> Research -> Drafting ⇄ Audit)..."):
        results = agent_system.invoke({"pdf_path": pdf_path, "loop_count": 0})
        st.session_state["results"] = results
        st.success("Agents Processing Completed!")

if "results" in st.session_state:
    res = st.session_state["results"]

    if not res.get("is_safe", True):
        st.error("🚨 Security Check Failed: File unsafe paayi gayi!")
    else:
        st.divider()
        st.header("[6] Human Review & Final DOCX Export")

        col1, col2 = st.columns(2)
        with col1:
            st.subheader("📋 Requirements Extracted")
            st.info(res.get("requirements", "N/A"))

            st.subheader("🔍 Research Data")
            st.write(res.get("research_data", "N/A"))
            st.caption(f"Audit Loops Completed: {res.get('loop_count', 0)}")

        with col2:
            st.subheader("✏️ Generated Draft (Editable)")
            user_edited_draft = st.text_area("Aap is draft ko edit kar sakte hain:", res.get("draft", ""), height=400)

            if st.button("Approve & Save Word Document"):
                doc = Document()
                doc.add_heading("Final Client Document", level=1)
                doc.add_paragraph(user_edited_draft)

                doc.save("Final_Output.docx")
                st.success("File 'Final_Output.docx' ke naam se save ho gayi hai!")
